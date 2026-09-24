import copy
import io
import json
from pathlib import Path
import shutil
import sys

import pytest
from PIL import Image
from pptx import Presentation

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / 'scripts'))
import exec_sketch as engine


@pytest.fixture
def job(tmp_path):
    shutil.copytree(SKILL / 'examples', tmp_path / 'job')
    return tmp_path / 'job/deck.json'


def edit(path, callback):
    value = json.loads(path.read_text(encoding='utf-8'))
    callback(value)
    engine.dump(path, value)


def test_bound_examples(job):
    assert engine.qa(job)['status'] == 'PASS'


def test_duplicate_id_rejected(job):
    edit(job, lambda d: d['slides'].append(copy.deepcopy(d['slides'][0])))
    with pytest.raises(ValueError, match='duplicate'):
        engine.load(job)


def test_unknown_theme_rejected(job):
    edit(job, lambda d: d['slides'][0].update(theme='unknown'))
    with pytest.raises(ValueError, match='Unknown theme'):
        engine.load(job)


@pytest.mark.parametrize('field,value', [('title', 'Changed'), ('visual', 'Changed diagram'),
                                       ('text', ['New text']), ('art_prompt', 'New prompt')])
def test_changed_prompt_requires_binding(job, field, value):
    edit(job, lambda d: d['slides'][0].update({field: value}))
    with pytest.raises(ValueError, match='Stale'):
        engine.qa(job)


def test_palette_change_invalidates(job):
    values = engine.themes()
    values['exec-sketch-light']['accent'] = '#FF0000'
    edit(job, lambda d: d.update(themes=values))
    with pytest.raises(ValueError, match='Stale'):
        engine.qa(job)


def test_image_hash_mismatch(job):
    deck = engine.load(job)
    path = job.parent / deck['slides'][0]['binding']['image']
    path.write_bytes(path.read_bytes() + b'changed')
    with pytest.raises(ValueError, match='Image hash'):
        engine.qa(job)


def test_missing_binding(job):
    edit(job, lambda d: d['slides'][0].pop('binding'))
    with pytest.raises(ValueError, match='Missing image'):
        engine.qa(job)


def test_true_magic_required(tmp_path):
    bad = tmp_path / 'bad.png'
    Image.new('RGB', (32, 32)).save(bad, format='JPEG')
    with pytest.raises(ValueError, match='extension'):
        engine.read_image(bad)


def test_no_traversal(tmp_path):
    with pytest.raises(ValueError, match='relative'):
        engine.safe_path(tmp_path, '../outside.png')


def test_installed_skill_is_not_a_job():
    with pytest.raises(ValueError, match='protected'):
        engine.writable(SKILL / 'examples/deck.json')


def test_contain_not_stretch():
    assert engine.contain((100, 50), (0, 0, 100, 100)) == (0, 25, 100, 50)


def test_bad_canvas(job):
    edit(job, lambda d: d.update(canvas={'width': 2858, 'height': 1524}))
    with pytest.raises(ValueError, match='16:9'):
        engine.load(job)


def test_paragraphs_not_hard_wrapped():
    assert engine.paragraphs(['Первая строка\nпродолжается.', 'Следующий абзац.']) == 'Первая строка продолжается.\n\nСледующий абзац.'


def test_notes_edit_keeps_art_binding(job):
    edit(job, lambda d: d['slides'][0]['notes'].update(spoken=['Новый полный абзац речи.']))
    assert engine.qa(job)['status'] == 'PASS'


def test_reorder_keeps_notes_with_id(job, tmp_path):
    edit(job, lambda d: d['slides'].reverse())
    out = tmp_path / 'out'
    assert engine.build(job, out)['status'] == 'PASS'
    deck = engine.load(job)
    pptx = Presentation(out / 'presentation.pptx')
    assert pptx.slides[0].notes_slide.notes_text_frame.text == engine.full_notes(deck['slides'][0])
    assert deck['slides'][0]['id'] == 'dark'
    assert abs(pptx.slide_width / pptx.slide_height - 16 / 9) < .00001
    assert len(pptx.slides) == 2
    assert (out / 'teleprompter.txt').read_text(encoding='utf-8-sig').count('РЕЖИССЁРСКИЕ') == 0
    assert 'https://' not in (out / 'preview.html').read_text(encoding='utf-8')


def test_no_output_clobber(job, tmp_path):
    out = tmp_path / 'out'
    out.mkdir()
    (out / 'existing.txt').write_text('keep')
    with pytest.raises(ValueError, match='new or empty'):
        engine.build(job, out)
    assert (out / 'existing.txt').read_text() == 'keep'


def test_plan_has_stable_ids(job, tmp_path):
    out = tmp_path / 'prompts'
    assert engine.plan(job, out)['planned'] == 2
    plan = json.loads((out / 'plan.json').read_text(encoding='utf-8'))
    assert [s['id'] for s in plan['slides']] == ['light', 'dark']
    assert all((out / s['file']).is_file() for s in plan['slides'])


def test_bind_strips_metadata_and_rejects_unknown(job, tmp_path):
    with pytest.raises(ValueError, match='Unknown slide'):
        engine.bind(job, 'missing', SKILL / 'examples/media/light.png', 'fixture')
    result = engine.bind(job, 'light', SKILL / 'examples/media/light.png', 'fixture')
    assert result['bound'] == 'light'
    assert engine.qa(job)['status'] == 'PASS'


def test_overlay_is_separate_and_contains(job, tmp_path):
    logo = job.parent / 'media/brand.png'
    Image.new('RGBA', (100, 50), '#3366FF').save(logo)
    item = {'role': 'logo', 'image': 'media/brand.png', 'sha256': engine.sha(logo.read_bytes()), 'rect': [.88, .02, .08, .08]}
    edit(job, lambda d: d.update(branding=[item]))
    for sid in ['light', 'dark']:
        engine.bind(job, sid, SKILL / f'examples/media/{sid}.png', 'fixture')
    out = tmp_path / 'with-overlay'
    engine.build(job, out)
    pptx = Presentation(out / 'presentation.pptx')
    assert len(pptx.slides[0].shapes) == 2
    shape = pptx.slides[0].shapes[1]
    assert abs(shape.width / shape.height - 2) < .00001
    assert engine.qa(job, out)['status'] == 'PASS'


def test_overlay_outside_frame(job):
    deck = engine.load(job)
    deck['branding'] = [{'role': 'logo', 'rect': [.99, 0, .1, .1]}]
    with pytest.raises(ValueError, match='beyond'):
        engine.overlays(deck, deck['slides'][0], job.parent)


def test_art_inset_letterboxes_with_artwork_paper(job, tmp_path):
    edit(job, lambda d: d.update(art_inset=.9))
    assert engine.qa(job)['status'] == 'PASS'          # layout-only: the art binding stays valid
    out = tmp_path / 'inset'
    engine.build(job, out)
    deck = engine.load(job)
    art = engine.read_image(job.parent / deck['slides'][0]['binding']['image'])
    token = engine.palette(deck, deck['slides'][0])['background']
    paper = engine.paper_color(art, token)
    render = Image.open(out / 'slides/01-light.png').convert('RGB')
    assert '#%02X%02X%02X' % render.getpixel((10, 10)) == paper       # margin continues the paper
    assert paper != token
    pptx = Presentation(out / 'presentation.pptx')
    assert str(pptx.slides[0].background.fill.fore_color.rgb) == paper[1:]   # no seam in PowerPoint either
    picture = pptx.slides[0].shapes[0]
    assert abs(picture.width / pptx.slide_width - .9) < .002
    assert abs(picture.left / pptx.slide_width - .05) < .002
    records = json.loads((out / 'build-manifest.json').read_text(encoding='utf-8'))['slides']
    assert [r['art_inset'] for r in records] == [.9, .9]
    assert engine.qa(job, out)['status'] == 'PASS'


def test_inset_on_transparent_art_keeps_palette_paper(job, tmp_path):
    art = Image.new('RGBA', (1920, 1080), (0, 0, 0, 0))              # transparent page, black RGB underneath
    art.paste((20, 20, 20, 255), (400, 400, 1500, 700))
    path = job.parent / 'media/transparent.png'
    art.save(path)
    engine.bind(job, 'light', path, 'fixture')
    edit(job, lambda d: d['slides'][0].update(art_inset=.9))
    out = tmp_path / 'rgba'
    engine.build(job, out)
    token = engine.palette(engine.load(job), engine.load(job)['slides'][0])['background']
    render = Image.open(out / 'slides/01-light.png').convert('RGB')
    assert '#%02X%02X%02X' % render.getpixel((10, 10)) == token.upper()   # was black before the fix


def test_qa_detects_export_built_with_other_inset(job, tmp_path):
    out = tmp_path / 'export'
    engine.build(job, out)
    edit(job, lambda d: d.update(art_inset=.9))
    assert engine.qa(job)['status'] == 'PASS'           # the art binding is still valid...
    with pytest.raises(ValueError, match='Export is stale'):
        engine.qa(job, out)                             # ...but that export no longer matches the manifest


def test_slide_inset_overrides_deck(job, tmp_path):
    edit(job, lambda d: (d.update(art_inset=.9), d['slides'][1].update(art_inset=1)))
    out = tmp_path / 'mixed'
    engine.build(job, out)
    records = json.loads((out / 'build-manifest.json').read_text(encoding='utf-8'))['slides']
    assert [r['art_inset'] for r in records] == [.9, 1.0]


@pytest.mark.parametrize('value', [.5, 1.2, True, '0.9'])
def test_art_inset_bad_values_rejected(job, value):
    edit(job, lambda d: d.update(art_inset=value))
    with pytest.raises(ValueError, match='art_inset'):
        engine.load(job)


def test_art_inset_with_overlay_rejected(job):
    logo = job.parent / 'media/brand.png'
    Image.new('RGBA', (100, 50), '#3366FF').save(logo)
    item = {'role': 'logo', 'image': 'media/brand.png', 'sha256': engine.sha(logo.read_bytes()), 'rect': [.88, .02, .08, .08]}
    edit(job, lambda d: d.update(branding=[item], art_inset=.9))
    with pytest.raises(ValueError, match='reserved areas'):
        engine.load(job)


def test_improved_example_builds_with_inset(tmp_path):
    shutil.copytree(SKILL / 'examples/improved', tmp_path / 'improved')
    deck = tmp_path / 'improved/deck.json'
    out = tmp_path / 'out'
    assert engine.build(deck, out)['status'] == 'PASS'
    records = json.loads((out / 'build-manifest.json').read_text(encoding='utf-8'))['slides']
    assert [r['id'] for r in records] == ['light-v4', 'dark-v4']
    assert all(r['art_inset'] == .92 for r in records)


def test_template_plans_with_semantic_prefix(tmp_path):
    shutil.copy2(SKILL / 'templates/deck.json', tmp_path / 'deck.json')
    assert engine.plan(tmp_path / 'deck.json', tmp_path / 'prompts')['planned'] == 1
    prompt = (tmp_path / 'prompts/first-slide.txt').read_text(encoding='utf-8')
    assert 'ACCENT BUDGET' in prompt and 'exec-sketch' in prompt and '#F7F3E8' in prompt


@pytest.fixture
def bridge(monkeypatch):
    import exec_generate as bridge
    calls = []

    def fake_call(model, prompt):
        calls.append(prompt)
        return Image.new('RGB', (1920, 1088), '#F7F3E8'), {'usage': None}
    monkeypatch.setattr(bridge, 'call', fake_call)
    monkeypatch.setattr(bridge, 'keys', lambda: None)
    bridge.calls = calls
    return bridge


def unbound(job):
    edit(job, lambda d: [s.pop('binding') for s in d['slides']])


def test_bridge_normalize():
    import exec_generate as bridge
    assert bridge.normalize(Image.new('RGB', (1920, 1088))).size == (1920, 1080)
    assert bridge.normalize(Image.new('RGB', (1344, 768))).size == (1344, 768)   # letterboxed later, never stretched


def test_bridge_generates_binds_and_reuses_with_true_provenance(job, bridge):
    unbound(job)
    assert bridge.main([str(job), '--model', 'gpt-image-2.5-sunburst']) == 0
    assert len(bridge.calls) == 2
    assert engine.qa(job)['status'] == 'PASS'
    unbound(job)                                        # same prompts again, different --model
    assert bridge.main([str(job), '--model', 'gemini-3.1-flash-image-preview']) == 0
    assert len(bridge.calls) == 2                       # reused, not paid twice
    assert {s['binding']['provider'] for s in engine.load(job)['slides']} == {'gpt-image-2.5-sunburst'}


def test_bridge_uses_current_prompt_not_stale_plan(job, bridge, tmp_path):
    unbound(job)
    edit(job, lambda d: [s.pop('art_prompt') for s in d['slides']])
    engine.plan(job, tmp_path / 'prompts')
    edit(job, lambda d: d['slides'][0].update(title='Новый заголовок'))
    with pytest.raises(SystemExit, match='stale'):
        bridge.main([str(job), '--prompts', str(tmp_path / 'prompts')])
    assert bridge.calls == []                           # refused before paying
    assert bridge.main([str(job), '--ids', 'light']) == 0
    assert 'Новый заголовок' in bridge.calls[0]
    assert engine.load(job)['slides'][0]['binding']['art_signature'] == engine.signature(engine.load(job), engine.load(job)['slides'][0])


def test_bridge_redraws_stale_slide_and_keeps_old_art(job, bridge):
    unbound(job)
    bridge.main([str(job)])
    edit(job, lambda d: d['slides'][0].update(art_prompt='A different drawing'))
    assert bridge.main([str(job)]) == 0                 # default picks the stale slide only
    assert len(bridge.calls) == 3 and bridge.calls[-1] == 'A different drawing'
    assert len(list((job.parent / 'art/superseded').glob('light-*.png'))) == 1
    assert engine.qa(job)['status'] == 'PASS'


def test_bridge_refuses_hand_placed_art_and_bad_ids(job, bridge):
    unbound(job)
    (job.parent / 'art').mkdir()
    Image.new('RGB', (1920, 1080)).save(job.parent / 'art/light.png')   # no sidecar: provenance unknown
    with pytest.raises(SystemExit, match='no provenance'):
        bridge.main([str(job), '--ids', 'light'])
    assert (job.parent / 'art/light.png').is_file()     # left in place, nothing paid
    with pytest.raises(SystemExit, match='empty'):
        bridge.main([str(job), '--ids', ' , '])
    with pytest.raises(SystemExit, match='Unknown'):
        bridge.main([str(job), '--ids', 'nope'])
    assert bridge.calls == []


def test_bridge_dedupes_ids_and_no_bind_leaves_manifest(job, bridge):
    unbound(job)
    before = job.read_bytes()
    assert bridge.main([str(job), '--ids', 'dark, dark ,dark', '--no-bind']) == 0
    assert len(bridge.calls) == 1
    assert job.read_bytes() == before


def test_corrupt_export_notes_detected(job, tmp_path):
    out = tmp_path / 'out'
    engine.build(job, out)
    pptx = Presentation(out / 'presentation.pptx')
    pptx.slides[0].notes_slide.notes_text_frame.text = 'Wrong notes'
    pptx.save(out / 'presentation.pptx')
    with pytest.raises(ValueError, match='notes mismatch'):
        engine.qa(job, out)
