"""Portable, offline-first exec-sketch prompt, binding, export and QA pipeline."""
from __future__ import annotations

import argparse
import base64
import hashlib
import html
import io
import json
import math
from pathlib import Path
import re
import sys

from PIL import Image, ImageOps, ImageDraw

SKILL = Path(__file__).resolve().parents[1]
ID_RE = re.compile(r'^[a-z0-9][a-z0-9_-]{0,63}$')
COLORS = ('background', 'ink', 'accent', 'secondary')


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def writable(path: Path) -> Path:
    path = path.resolve()
    protected = [SKILL, *(Path.home() / name for name in ('.agents', '.claude', '.codex')), Path.home() / 'CLAUDE.md']
    if any(path == root.resolve() or path.is_relative_to(root.resolve()) for root in protected):
        raise ValueError('Write target is inside a protected source or installed skill; use a separate job')
    return path


def dump(path: Path, value) -> None:
    writable(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    temp.replace(path)


def safe_path(root: Path, value: str) -> Path:
    path = (root / value).resolve()
    if Path(value).is_absolute() or not path.is_relative_to(root.resolve()):
        raise ValueError('Asset paths must be relative and remain inside the project')
    return path


def read_image(path: Path) -> Image.Image:
    expected = {'.png': 'PNG', '.jpg': 'JPEG', '.jpeg': 'JPEG', '.webp': 'WEBP'}
    with Image.open(path) as image:
        if expected.get(path.suffix.lower()) != image.format:
            raise ValueError('Image extension does not match its verified format')
        image.load()
        return ImageOps.exif_transpose(image).convert('RGBA')


def themes() -> dict:
    return json.loads((SKILL / 'assets/palettes.json').read_text(encoding='utf-8'))


def palette(deck: dict, slide: dict) -> dict:
    values = themes() | deck.get('themes', {})
    if slide.get('theme') not in values:
        raise ValueError('Unknown theme: ' + str(slide.get('theme')))
    result = values[slide['theme']]
    for key in COLORS:
        if not re.fullmatch(r'#[0-9a-fA-F]{6}', result.get(key, '')):
            raise ValueError('Palette requires valid six-digit hex colors: ' + key)
    return result


def paragraphs(value) -> str:
    if not isinstance(value, list) or not value or any(not isinstance(p, str) or not p.strip() for p in value):
        raise ValueError('Spoken notes must be a non-empty array of paragraphs')
    return '\n\n'.join(' '.join(p.split()) for p in value)


def load(path: Path) -> dict:
    deck = json.loads(path.read_text(encoding='utf-8-sig'))
    if deck.get('version') != 1:
        raise ValueError('Unsupported manifest version')
    canvas = deck.get('canvas', {})
    w, h = canvas.get('width', 0), canvas.get('height', 0)
    if not isinstance(w, int) or not isinstance(h, int) or w < 640 or h < 360 or w > 7680 or h > 4320 or w * 9 != h * 16:
        raise ValueError('Canvas must be 16:9, between 640x360 and 7680x4320')
    slides = deck.get('slides', [])
    if not slides or len(slides) > 200:
        raise ValueError('Provide between 1 and 200 slides')
    ids = set()
    for slide in slides:
        sid = slide.get('id', '')
        if not ID_RE.fullmatch(sid) or sid in ids:
            raise ValueError('Invalid or duplicate slide ID: ' + sid)
        ids.add(sid)
        for key in ('title', 'visual'):
            if not isinstance(slide.get(key), str) or not slide[key].strip():
                raise ValueError(sid + ' requires ' + key)
        if not isinstance(slide.get('text'), list) or not all(isinstance(x, str) for x in slide['text']):
            raise ValueError(sid + ' requires an exact text array')
        paragraphs(slide.get('notes', {}).get('spoken'))
        palette(deck, slide)
        art_inset(deck, slide)
    return deck


def art_inset(deck: dict, slide: dict) -> float:
    """Build-time scale of the artwork, 0.8..1.0; 1.0 = full bleed. Not part of the art signature."""
    value = slide.get('art_inset', deck.get('art_inset', 1.0))
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not .8 <= value <= 1:
        raise ValueError('art_inset must be a number between 0.8 and 1.0')
    if value < 1 and (deck.get('branding') or slide.get('overlays')):
        raise ValueError('art_inset cannot be combined with branding or overlays: reserved areas would move')
    return float(value)


def paper_color(image: Image.Image, background: str) -> str:
    """Median of the four corners: generated paper drifts from the palette token.

    Transparent art is seen over the palette background, so it is sampled that way too;
    converting RGBA straight to RGB would read the black behind transparent pixels.
    """
    rgb = Image.alpha_composite(Image.new('RGBA', image.size, background), image.convert('RGBA')).convert('RGB')
    w, h = rgb.size
    samples = [rgb.getpixel((x, y)) for x in (2, w - 3) for y in (2, h - 3)]
    return '#%02X%02X%02X' % tuple(sorted(c[i] for c in samples)[len(samples) // 2] for i in range(3))


def make_prompt(deck: dict, slide: dict) -> str:
    if slide.get('art_prompt'):
        return slide['art_prompt']
    contract = (SKILL / 'references/prompts/semantic-prefix.txt').read_text(encoding='utf-8').strip()
    return '\n\n'.join([
        contract,
        'PALETTE AND TYPOGRAPHY:\n' + json.dumps(palette(deck, slide), ensure_ascii=False),
        'LANGUAGE: ' + deck.get('language', 'Russian'),
        'EXACT HEADING:\n' + slide['title'],
        'EXACT OTHER TEXT (each once):\n' + '\n'.join(slide['text']),
        'VISUAL ARGUMENT:\n' + slide['visual'],
        'OPTIONAL ASSET RESERVATIONS:\n' + json.dumps([
            {'role': o['role'], 'rect': o['rect']} for o in deck.get('branding', []) + slide.get('overlays', [])
        ], ensure_ascii=False)
    ])


def signature(deck: dict, slide: dict) -> str:
    value = {'prompt': make_prompt(deck, slide), 'palette': palette(deck, slide),
             'canvas': deck['canvas'], 'title': slide['title'], 'text': slide['text'],
             'visual': slide['visual'], 'reservations': [
                 {'role': o['role'], 'rect': o['rect']} for o in deck.get('branding', []) + slide.get('overlays', [])]}
    return sha(json.dumps(value, ensure_ascii=False, sort_keys=True).encode('utf-8'))


def plan(manifest: Path, output: Path) -> dict:
    deck = load(manifest)
    writable(output)
    output.mkdir(parents=True, exist_ok=True)
    rows = []
    for slide in deck['slides']:
        prompt = make_prompt(deck, slide)
        filename = slide['id'] + '.txt'
        (output / filename).write_text(prompt + '\n', encoding='utf-8')
        rows.append({'id': slide['id'], 'theme': slide['theme'], 'file': filename,
                     'prompt_sha256': sha(prompt.encode('utf-8')), 'art_signature': signature(deck, slide)})
    dump(output / 'plan.json', {'version': 1, 'slides': rows, 'generation': 'Not performed; use an authorized image tool'})
    return {'planned': len(rows)}


def bind(manifest: Path, slide_id: str, image_path: Path, provider: str) -> dict:
    writable(manifest)
    deck = load(manifest)
    slide = next((s for s in deck['slides'] if s['id'] == slide_id), None)
    if slide is None:
        raise ValueError('Unknown slide ID')
    image = read_image(image_path)
    # Re-encode to PNG to avoid EXIF metadata and misleading extensions.
    data = io.BytesIO()
    image.save(data, format='PNG')
    digest = sha(data.getvalue())
    rel = f'media/{slide_id}-{digest[:16]}.png'
    dest = safe_path(manifest.parent, rel)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data.getvalue())
    slide['binding'] = {'image': rel, 'image_sha256': digest,
                        'prompt_sha256': sha(make_prompt(deck, slide).encode('utf-8')),
                        'art_signature': signature(deck, slide), 'provider': provider}
    dump(manifest, deck)
    return {'bound': slide_id, 'image_sha256': digest}


def contain(size: tuple[int, int], box: tuple[float, float, float, float]) -> tuple[float, float, float, float]:
    x, y, w, h = box
    scale = min(w / size[0], h / size[1])
    iw, ih = size[0] * scale, size[1] * scale
    return x + (w - iw) / 2, y + (h - ih) / 2, iw, ih


def overlays(deck: dict, slide: dict, root: Path) -> list:
    result = []
    for item in deck.get('branding', []) + slide.get('overlays', []):
        rect = item.get('rect', [])
        if len(rect) != 4 or any(not isinstance(v, (int, float)) or not math.isfinite(v) for v in rect):
            raise ValueError('Overlay rect must be [x, y, width, height], normalized 0..1')
        x, y, w, h = rect
        if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 1.000001 or y + h > 1.000001:
            raise ValueError('Overlay extends beyond the slide')
        path = safe_path(root, item['image'])
        if not item.get('sha256') or sha(path.read_bytes()) != item['sha256']:
            raise ValueError('Overlay hash missing or changed')
        result.append((item, read_image(path)))
    return result


def validate_bound(deck: dict, root: Path) -> list:
    warnings = []
    for slide in deck['slides']:
        binding = slide.get('binding', {})
        if not binding:
            raise ValueError('Missing image binding: ' + slide['id'])
        if binding.get('art_signature') != signature(deck, slide):
            raise ValueError('Stale image after prompt/layout change: ' + slide['id'])
        if binding.get('prompt_sha256') != sha(make_prompt(deck, slide).encode('utf-8')):
            raise ValueError('Prompt hash mismatch: ' + slide['id'])
        path = safe_path(root, binding['image'])
        if sha(path.read_bytes()) != binding.get('image_sha256'):
            raise ValueError('Image hash mismatch: ' + slide['id'])
        image = read_image(path)
        if abs(image.width / image.height - 16 / 9) > .01:
            warnings.append(slide['id'] + ': artwork is not 16:9; letterboxing will preserve proportions')
        if len(paragraphs(slide['notes']['spoken']).split()) < 70:
            warnings.append(slide['id'] + ': short spoken notes; review against intended duration')
        overlays(deck, slide, root)
    return warnings


def full_notes(slide: dict) -> str:
    text = paragraphs(slide['notes']['spoken'])
    if slide['notes'].get('stage'):
        text += '\n\n[РЕЖИССЁРСКИЕ ПОМЕТКИ — НЕ ЧИТАТЬ]\n' + '\n'.join(slide['notes']['stage'])
    if slide.get('sources'):
        text += '\n\n[ИСТОЧНИКИ]\n' + '\n'.join(slide['sources'])
    return text


def png_bytes(image: Image.Image) -> bytes:
    buf = io.BytesIO()
    image.save(buf, format='PNG')
    return buf.getvalue()


def build(manifest: Path, output: Path) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    import fitz

    deck = load(manifest)
    warnings = validate_bound(deck, manifest.parent)
    writable(output)
    if output.exists() and any(output.iterdir()):
        raise ValueError('Build output must be a new or empty directory')
    output.mkdir(parents=True, exist_ok=True)
    (output / 'slides').mkdir(exist_ok=True)
    presentation = Presentation()
    presentation.core_properties.author = ''
    presentation.core_properties.last_modified_by = ''
    presentation.core_properties.title = deck.get('title', 'Exec-sketch presentation')
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    w, h = deck['canvas']['width'], deck['canvas']['height']
    pdf = fitz.open()
    notes, reading, html_pages, rendered, result_rows = [], [], [], [], []
    motion = {'version': 1, 'canvas': deck['canvas'], 'timing_status': 'Needs transcript alignment; no timings invented', 'slides': []}
    for index, slide in enumerate(deck['slides'], 1):
        artwork = read_image(safe_path(manifest.parent, slide['binding']['image']))
        inset = art_inset(deck, slide)
        # With an inset the margin must continue the artwork's own paper, or a seam shows.
        background = palette(deck, slide)['background']
        if inset < 1:
            background = paper_color(artwork, background)
        frame = Image.new('RGBA', (w, h), background)
        ppt_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        ppt_slide.background.fill.solid()
        ppt_slide.background.fill.fore_color.rgb = RGBColor.from_string(background[1:])
        layers = [(artwork, (w * (1 - inset) / 2, h * (1 - inset) / 2, w * inset, h * inset))]
        for item, image in overlays(deck, slide, manifest.parent):
            x, y, iw, ih = item['rect']
            layers.append((image, (x * w, y * h, iw * w, ih * h)))
        for image, box in layers:
            x, y, iw, ih = contain(image.size, box)
            frame.alpha_composite(image.resize((max(1, round(iw)), max(1, round(ih))), Image.Resampling.LANCZOS), (round(x), round(y)))
            ppt_slide.shapes.add_picture(io.BytesIO(png_bytes(image)), round(x / w * presentation.slide_width),
                                        round(y / h * presentation.slide_height),
                                        round(iw / w * presentation.slide_width), round(ih / h * presentation.slide_height))
        ppt_slide.notes_slide.notes_text_frame.text = full_notes(slide)
        image_path = output / 'slides' / f'{index:02d}-{slide["id"]}.png'
        frame.convert('RGB').save(image_path)
        rendered.append(frame.convert('RGB'))
        pdf.new_page(width=960, height=540).insert_image(fitz.Rect(0, 0, 960, 540), stream=image_path.read_bytes())
        header = f'{index:02d}. {slide["title"]}'
        notes.append(header + '\n\n' + full_notes(slide))
        reading.append(header + '\n\n' + paragraphs(slide['notes']['spoken']))
        html_pages.append('<section><img alt="' + html.escape(slide['title'], quote=True) + '" src="data:image/png;base64,' + base64.b64encode(image_path.read_bytes()).decode() + '"></section>')
        result_rows.append({'index': index, 'id': slide['id'], 'prompt_sha256': slide['binding']['prompt_sha256'],
                            'image_sha256': slide['binding']['image_sha256'], 'art_inset': inset,
                            'notes_sha256': sha(full_notes(slide).encode('utf-8')),
                            'render': image_path.relative_to(output).as_posix(), 'render_sha256': sha(image_path.read_bytes())})
        motion['slides'].append({'id': slide['id'], 'render': image_path.relative_to(output).as_posix(),
                                'theme': slide['theme'], 'ink': palette(deck, slide)['ink'], 'accent': palette(deck, slide)['accent'],
                                'draw_order': slide.get('draw_order', []), 'spoken': slide['notes']['spoken']})
    presentation.save(output / 'presentation.pptx')
    pdf.save(output / 'presentation.pdf', garbage=4, deflate=True)
    pdf.close()
    (output / 'speaker-notes.txt').write_text('\n\n———\n\n'.join(notes) + '\n', encoding='utf-8-sig')
    (output / 'teleprompter.txt').write_text('\n\n———\n\n'.join(reading) + '\n', encoding='utf-8-sig')
    (output / 'preview.html').write_text('<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Exec-sketch preview</title><style>body{margin:0;background:#25272a}section{max-width:1280px;margin:24px auto}img{display:block;width:100%;height:auto}</style>' + ''.join(html_pages) + '</html>', encoding='utf-8')
    contact = Image.new('RGB', (960, math.ceil(len(rendered) / 2) * 290), '#dedede')
    draw = ImageDraw.Draw(contact)
    for i, image in enumerate(rendered):
        contact.paste(image.resize((480, 270), Image.Resampling.LANCZOS), ((i % 2) * 480, (i // 2) * 290))
        draw.text(((i % 2) * 480 + 10, (i // 2) * 290 + 271), f'{i+1:02d} / {deck["slides"][i]["id"]}', fill='black')
    contact.save(output / 'contact-sheet.png')
    dump(output / 'motion-handoff.json', motion)
    dump(output / 'build-manifest.json', {'version': 1, 'slides': result_rows, 'warnings': warnings,
                                        'visual_qa': 'Not automated: inspect all slides and PowerPoint rendering'})
    report = qa(manifest, output)
    dump(output / 'qa.json', report)
    return report


def qa(manifest: Path, output: Path | None = None) -> dict:
    from pptx import Presentation
    import fitz

    deck = load(manifest)
    warnings = validate_bound(deck, manifest.parent)
    if output:
        presentation = Presentation(output / 'presentation.pptx')
        records = json.loads((output / 'build-manifest.json').read_text(encoding='utf-8'))['slides']
        if len(presentation.slides) != len(deck['slides']) or len(records) != len(deck['slides']):
            raise ValueError('PPTX/build slide count mismatch')
        for index, (original, exported, record) in enumerate(zip(deck['slides'], presentation.slides, records)):
            if record['id'] != original['id'] or record['notes_sha256'] != sha(full_notes(original).encode('utf-8')):
                raise ValueError('Build manifest order or notes mismatch')
            if (record['image_sha256'] != original['binding']['image_sha256']
                    or record['prompt_sha256'] != original['binding']['prompt_sha256']
                    or record.get('art_inset', 1.0) != art_inset(deck, original)):
                raise ValueError('Export is stale for ' + original['id'] + ': rebuild into a new directory')
            if exported.notes_slide.notes_text_frame.text != full_notes(original):
                raise ValueError('PPTX notes mismatch at slide ' + str(index + 1))
            if len(exported.shapes) != 1 + len(deck.get('branding', [])) + len(original.get('overlays', [])):
                raise ValueError('PPTX layer count mismatch')
            render = safe_path(output, record['render'])
            if sha(render.read_bytes()) != record['render_sha256']:
                raise ValueError('Rendered image changed')
        with fitz.open(output / 'presentation.pdf') as pdf:
            if pdf.page_count != len(deck['slides']):
                raise ValueError('PDF slide count mismatch')
    return {'status': 'PASS', 'slides': len(deck['slides']), 'warnings': warnings,
            'scope': 'Technical manifest, hashes, image formats, aspect ratio policy, overlay boxes; exports and notes when provided',
            'requires_human_review': ['Cyrillic and facts inside generated art', 'Visual density and collisions', 'Actual PowerPoint application rendering']}


def utf8_stdio() -> None:
    """Windows consoles default to cp1251/cp866/cp1252; JSON with Cyrillic must not crash the CLI."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8', errors='replace')
        except Exception:
            pass


def main() -> int:
    utf8_stdio()
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    for cmd in ('plan', 'build', 'qa'):
        p = commands.add_parser(cmd)
        p.add_argument('manifest', type=Path)
        p.add_argument('--out', type=Path, required=cmd != 'qa')
    p = commands.add_parser('bind')
    p.add_argument('manifest', type=Path)
    p.add_argument('--id', required=True)
    p.add_argument('--image', type=Path, required=True)
    p.add_argument('--provider', required=True)
    args = parser.parse_args()
    try:
        if args.command == 'bind':
            result = bind(args.manifest.resolve(), args.id, args.image.resolve(), args.provider)
        else:
            result = globals()[args.command](args.manifest.resolve(), args.out.resolve() if args.out else None)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ValueError, OSError, KeyError, TypeError) as error:
        # Do not print exception payloads that could contain local paths or secrets.
        print(json.dumps({'status': 'FAIL', 'error_type': type(error).__name__,
                          'message': str(error) if isinstance(error, ValueError) else 'Invalid input or unavailable file; check the manifest'}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    sys.exit(main())
